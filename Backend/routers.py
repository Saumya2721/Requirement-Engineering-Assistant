import json
import io
import traceback
import logging
from typing import Dict, Any
from fastapi import APIRouter, HTTPException, Request, Depends
from fastapi.responses import StreamingResponse
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

from schemas import TurnRequest, SynthesisRequest, PipelineState
from services import ambiguity_chain, requirement_graph
from dependencies import verify_api_key
from limiter import limiter

router = APIRouter(prefix="/api")
logger = logging.getLogger(__name__)
logging.basicConfig(level=logging.INFO)

@router.post("/analyze-turn")
@limiter.limit("60/minute")
async def analyze_turn(request: Request, req: TurnRequest, api_key: str = Depends(verify_api_key)):
    try:
        result = await ambiguity_chain.ainvoke({
            "context": req.context,
            "latest_statement": req.latest_statement
        })
        return result
    except Exception as e:
        logger.warning("Analyze-turn LLM/RateLimit warning (graceful fallback): %s", str(e))
        return {
            "has_ambiguity": False,
            "ambiguous_phrases": [],
            "clarification_question": None
        }

@router.post("/synthesize")
@limiter.limit("5/minute")
async def synthesize_requirements(request: Request, req: SynthesisRequest, api_key: str = Depends(verify_api_key)):
    try:
        initial_state: PipelineState = {
            "raw_transcript": req.raw_transcript,
            "clarifications": req.clarifications,
            "baseline_fr": [],
            "baseline_nfr": {},
            "refined_fr": [],
            "refined_nfr": {},
            "evaluation_report": {}
        }
        final_state = await requirement_graph.ainvoke(initial_state)
        return final_state
    except Exception as e:
        logger.error("Error synthesizing requirements: %s", e)
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/export")
@limiter.limit("10/minute")
async def export_document(request: Request, payload: Dict[str, Any], api_key: str = Depends(verify_api_key)):
    export_format = payload.get("format", "txt").lower()
    data = payload.get("data", {})
    
    if export_format == "txt":
        content = f"=== SOFTWARE REQUIREMENT SPECIFICATION & EVALUATION ===\n\n"
        content += f"--- BASELINE REQUIREMENTS (Without Clarification) ---\n"
        content += f"Functional Requirements:\n" + "\n".join([f"- {r}" for r in data.get("baseline_fr", [])]) + "\n\n"
        content += f"Non-Functional Requirements:\n"
        for cat, reqs in data.get("baseline_nfr", {}).items():
            content += f"[{cat}]\n" + "\n".join([f"  * {r}" for r in reqs]) + "\n"
            
        content += f"\n--- REFINED REQUIREMENTS (With Clarification) ---\n"
        content += f"Functional Requirements:\n" + "\n".join([f"- {r}" for r in data.get("refined_fr", [])]) + "\n\n"
        content += f"Non-Functional Requirements:\n"
        for cat, reqs in data.get("refined_nfr", {}).items():
            content += f"[{cat}]\n" + "\n".join([f"  * {r}" for r in reqs]) + "\n"
            
        content += f"\n--- EVALUATION & COMPARISON REPORT ---\n"
        metrics = data.get("evaluation_report", {}).get("metrics", {})
        content += f"Baseline Scores: {json.dumps(metrics.get('baseline', {}), indent=2)}\n"
        content += f"Refined Scores:  {json.dumps(metrics.get('refined', {}), indent=2)}\n"
        content += f"\nKey Improvements:\n" + "\n".join([f"- {i}" for i in data.get("evaluation_report", {}).get("key_improvements", [])])
        
        return StreamingResponse(
            io.BytesIO(content.encode("utf-8")),
            media_type="text/plain",
            headers={"Content-Disposition": "attachment; filename=Requirements_Report.txt"}
        )

    elif export_format == "docx":
        doc = Document()
        doc.add_heading("Software Requirements Analysis & Evaluation Report", level=0)
        
        doc.add_heading("1. Refined Requirements (Post-Clarification)", level=1)
        doc.add_heading("Functional Requirements", level=2)
        for fr in data.get("refined_fr", []):
            doc.add_paragraph(fr, style='List Bullet')
            
        doc.add_heading("Non-Functional Requirements", level=2)
        for cat, reqs in data.get("refined_nfr", {}).items():
            doc.add_heading(cat, level=3)
            for r in reqs:
                doc.add_paragraph(r, style='List Bullet')
                
        doc.add_heading("2. Quality Evaluation & Comparison", level=1)
        doc.add_paragraph(data.get("evaluation_report", {}).get("summary_comparison", ""))
        
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        return StreamingResponse(
            file_stream,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": "attachment; filename=Requirements_Report.docx"}
        )

    elif export_format == "pdf":
        file_stream = io.BytesIO()
        pdf = SimpleDocTemplate(file_stream, pagesize=letter)
        styles = getSampleStyleSheet()
        elements = [
            Paragraph("<b>Software Requirements & Evaluation Report</b>", styles['Title']),
            Spacer(1, 12),
            Paragraph("<b>Refined Functional Requirements</b>", styles['Heading2'])
        ]
        for fr in data.get("refined_fr", []):
            elements.append(Paragraph(f"• {fr}", styles['BodyText']))
            
        elements.append(Spacer(1, 10))
        elements.append(Paragraph("<b>Refined Non-Functional Requirements</b>", styles['Heading2']))
        for cat, reqs in data.get("refined_nfr", {}).items():
            elements.append(Paragraph(f"<b>{cat}:</b>", styles['Heading3']))
            for r in reqs:
                elements.append(Paragraph(f"• {r}", styles['BodyText']))
                
        elements.append(Spacer(1, 10))
        elements.append(Paragraph("<b>Quality Comparison & Summary</b>", styles['Heading2']))
        elements.append(Paragraph(data.get("evaluation_report", {}).get("summary_comparison", ""), styles['BodyText']))
        
        pdf.build(elements)
        file_stream.seek(0)
        return StreamingResponse(
            file_stream,
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename=Requirements_Report.pdf"}
        )
